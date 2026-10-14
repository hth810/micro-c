from pathlib import Path
import zipfile

from docx import Document
from docx.oxml.ns import qn
import pytest

from scripts.update_final_report_20261014 import build, page_segments


def reference(path: Path):
    doc = Document()
    doc.add_heading("Original report", 0)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "field"
    table.cell(0, 1).text = "value"
    table.cell(1, 0).text = "old"
    table.cell(1, 1).text = "1"
    p = doc.add_heading("Second page", 1)
    p.paragraph_format.page_break_before = True
    doc.add_paragraph("Untouched content")
    doc.save(path)


def test_edits_preserve_opaque_reference_parts(tmp_path):
    source, final = tmp_path / "source.docx", tmp_path / "final.docx"
    reference(source)
    result = build(source, {"reference_pages": 2, "pages": {"1": [
        {"type": "title", "text": "Final report"},
        {"type": "table", "caption": "Table 1", "headers": ["field", "value"],
         "rows": [["new", "2"], ["other", "3"]], "widths": [2, 4.86]},
    ]}}, final)
    doc = Document(final)
    assert doc.paragraphs[0].text == "Final report"
    assert doc.paragraphs[1].text == "表 1  Table 1"
    assert result["caption_counts"] == {"表": 1, "图": 0}
    assert doc.tables[0].cell(2, 1).text == "3"
    assert doc.paragraphs[-1].text == "Untouched content"
    assert len(page_segments(doc._element.body)) == 2
    assert not doc.paragraphs[0].runs[0]._element.xpath("./w:rPr/w:b")
    with zipfile.ZipFile(source) as a, zipfile.ZipFile(final) as b:
        for name in result["preserved_package_parts"]:
            assert a.read(name) == b.read(name)
        assert a.read("word/styles.xml") == b.read("word/styles.xml")
    assert doc._element.body[-1].tag == qn("w:sectPr")


def test_no_overwrite_and_slot_validation(tmp_path):
    source, final = tmp_path / "source.docx", tmp_path / "final.docx"
    reference(source)
    with pytest.raises(FileExistsError):
        build(source, {}, source)
    with pytest.raises(ValueError, match="pagination"):
        build(source, {"reference_pages": 26}, final)
    with pytest.raises(ValueError, match="Unknown"):
        build(source, {"reference_pages": 2, "pages": {"4": []}}, final)
