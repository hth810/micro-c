"""Edit selected report pages while preserving the reference package and style.

Patch content is evidence-backed JSON, supplied only after experiment results
have been reviewed. Existing report and historical receipts are never changed.
On Windows without the supplied document renderer, use the isolated Word export
helper, then inspect every rendered page. No external Python dependencies are
installed by this script.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
FONT = "Microsoft YaHei"
WIDTH = 6.86


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def font(run, *, size=None, bold=None, color="000000"):
    run.font.name = FONT
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    run.italic = False
    if size is not None:
        run.font.size = Pt(size)
    fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().insert(0, fonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), FONT)
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + name), None)


def paragraph(doc, text, style="Normal", *, first=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.widow_control = True
    if first:
        p.paragraph_format.page_break_before = True
    font(p.add_run(text), size=9.5 if style == "Caption" else None)
    if style == "Caption":
        p.paragraph_format.keep_together = True
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(7)
    return p


def table_from_template(doc, block, templates):
    headers, rows = block["headers"], block["rows"]
    widths = block["widths"]
    if len(headers) != len(widths) or abs(sum(widths) - WIDTH) > .01:
        raise ValueError("Explicit table widths must equal reference body width")
    if any(len(row) != len(headers) for row in rows):
        raise ValueError("Incomplete report table row")
    template = templates.get(len(headers))
    if template is None:
        raise ValueError("No reference table component for this column count")
    element = deepcopy(template)
    old_rows = element.findall(qn("w:tr"))
    header, body = deepcopy(old_rows[0]), deepcopy(old_rows[1])
    for row in old_rows:
        element.remove(row)
    grid = element.find(qn("w:tblGrid"))
    for cell, width in zip(grid, widths):
        cell.set(qn("w:w"), str(round(width * 1440)))
    # Add the cloned component, then use python-docx only for text/width edits.
    doc._element.body.insert(len(doc._element.body) - 1, element)
    from docx.table import Table
    table = Table(element, doc._body)
    for rid, values in enumerate([headers] + rows):
        element.append(deepcopy(header if rid == 0 else body))
        row = table.rows[-1]
        for cell, width, text in zip(row.cells, widths, values):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = str(text)
            shading = cell._tc.get_or_add_tcPr().find(qn("w:shd"))
            if shading is not None:
                shading.set(qn("w:fill"), "173B50" if rid == 0 else ("F1F5F8" if rid % 2 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(14)
            short = rid == 0 or bool(re.fullmatch(r"[0-9\s./%+\-−–→\[\],]+", str(text)))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if short else WD_ALIGN_PARAGRAPH.LEFT
            font(p.runs[0], size=10, bold=rid == 0, color="FFFFFF" if rid == 0 else "000000")
    return element


def page_segments(body):
    items = list(body)
    items = [item for item in items if item.tag != qn("w:sectPr")]
    starts = [0]
    for index, item in enumerate(items):
        flag = item.find("w:pPr/w:pageBreakBefore", item.nsmap)
        if index and flag is not None and flag.get(qn("w:val"), "1") not in ("0", "false"):
            starts.append(index)
    starts.append(len(items))
    return {i + 1: items[a:b] for i, (a, b) in enumerate(zip(starts, starts[1:]))}


def build(source: Path, patch: dict, destination: Path) -> dict:
    if source.resolve() == destination.resolve() or destination.exists():
        raise FileExistsError("Preserve reference and previous final files; choose a new output")
    original = source.read_bytes()
    doc = Document(io.BytesIO(original))
    segments = page_segments(doc._element.body)
    if len(segments) != patch.get("reference_pages", 26):
        raise ValueError("Reference pagination slots differ; inspect before editing")
    templates = {len(t.columns): deepcopy(t._tbl) for t in doc.tables}
    sources = []
    for number, blocks in sorted(patch["pages"].items(), key=lambda item: int(item[0])):
        number = int(number)
        if number not in segments or not blocks:
            raise ValueError("Unknown or empty page slot")
        old = segments[number]
        anchor = old[0]
        index = list(doc._element.body).index(anchor)
        fresh = []
        for pos, block in enumerate(blocks):
            before = set(doc._element.body)
            kind = block["type"]
            if kind in ("p", "h1", "h2", "small", "title", "subtitle"):
                style = {"h1": "Heading 1", "h2": "Heading 2", "title": "Title", "subtitle": "Subtitle"}.get(kind, "Normal")
                p = paragraph(doc, block["text"], style, first=bool((number > 1 and pos == 0) or block.get("new_page")))
                if kind == "small":
                    font(p.runs[0], size=9.5)
            elif kind == "table":
                paragraph(doc, block["caption"], "Caption", first=bool((number > 1 and pos == 0) or block.get("new_page")))
                table_from_template(doc, block, templates)
            elif kind == "figure":
                path = Path(block["path"])
                path = path if path.is_absolute() else ROOT / path
                data = path.read_bytes()
                iw, ih = Image.open(io.BytesIO(data)).size
                width = min(block.get("width", WIDTH), WIDTH)
                height = block.get("max_height", 3.0)
                width = min(width, height * iw / ih)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(12)
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
                shape = p.add_run().add_picture(io.BytesIO(data), width=Inches(width))
                shape._inline.docPr.set("descr", block["caption"])
                paragraph(doc, block["caption"], "Caption")
                sources.append({"page_slot": number, "path": str(path), "sha256": sha(data), "caption": block["caption"]})
            else:
                raise ValueError(f"Unsupported edit component {kind}")
            fresh.extend(item for item in doc._element.body if item not in before)
        for item in old:
            doc._element.body.remove(item)
        for item in fresh:
            doc._element.body.remove(item)
        for offset, item in enumerate(fresh):
            doc._element.body.insert(index + offset, item)
    replacements = patch.get("language_replacements", {})
    for p in doc.paragraphs:
        text = p.text
        for old, new in replacements.items():
            text = text.replace(old, new)
        if text != p.text and p.runs:
            p.runs[0].text = text
            for run in p.runs[1:]:
                run.text = ""
    # Renumber actual caption components in final reading order. Added evidence
    # tables cannot leave duplicate or out-of-order caption numbers.
    counters = {"表": 0, "图": 0}
    for p in doc.paragraphs:
        if p.style.name != "Caption":
            continue
        match = re.match(r"^(表|图)\s*\d+\s*", p.text)
        key = "表" if p._p.getnext() is not None and p._p.getnext().tag == qn("w:tbl") else "图"
        counters[key] += 1
        text = f"{key} {counters[key]}  " + (p.text[match.end():] if match else p.text)
        if p.runs:
            p.runs[0].text = text
            for run in p.runs[1:]:
                run.text = ""
    # Merge only editable package parts. Every untouched original package part
    # (styles, themes, headers, footers, old images, numbering) remains byte exact.
    generated = io.BytesIO()
    doc.save(generated)
    allowed = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    preserved = {}
    destination.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(io.BytesIO(original)) as reference, zipfile.ZipFile(generated) as changed, zipfile.ZipFile(destination, "x", zipfile.ZIP_DEFLATED) as out:
        prior = set(reference.namelist())
        for info in reference.infolist():
            data = changed.read(info.filename) if info.filename in allowed else reference.read(info.filename)
            out.writestr(info, data)
            if info.filename not in allowed:
                preserved[info.filename] = sha(data)
        for info in changed.infolist():
            if info.filename not in prior:
                if not info.filename.startswith("word/media/"):
                    raise AssertionError(f"Unexpected new package part {info.filename}")
                out.writestr(info, changed.read(info.filename))
    with zipfile.ZipFile(destination) as final:
        if any(sha(final.read(name)) != value for name, value in preserved.items()):
            raise AssertionError("Reference preserve-only package regression")
    return {"completed_at_utc": datetime.now(timezone.utc).isoformat(),
            "reference": str(source), "reference_sha256": sha(original),
            "output": str(destination), "output_sha256": sha(destination.read_bytes()),
            "edited_page_slots": sorted(map(int, patch["pages"])), "caption_counts": counters,
            "preserved_package_parts": preserved, "figure_sources": sources,
            "visual_status": "not accepted until all final page renders are inspected"}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--source", type=Path, required=True)
    p.add_argument("--patch", type=Path, required=True)
    p.add_argument("--out", type=Path, required=True)
    p.add_argument("--receipt", type=Path, required=True)
    args = p.parse_args()
    if args.receipt.exists():
        raise FileExistsError("Preserve previous report receipt")
    patch = json.loads(args.patch.read_text(encoding="utf-8"))
    result = build(args.source, patch, args.out)
    args.receipt.parent.mkdir(parents=True, exist_ok=True)
    args.receipt.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k not in ("preserved_package_parts", "figure_sources")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
