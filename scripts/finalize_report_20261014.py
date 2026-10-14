"""Check the exact final report bytes and record completed visual review.

The reviewed-pages argument records an actual human/agent page inspection;
the structural checks do not infer visual quality or scientific success.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
SUPPORT = ROOT / "output/report/support_final_verification_20261014"


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render-dir", type=Path, required=True)
    parser.add_argument("--reviewed-pages", required=True)
    args = parser.parse_args()
    destination = ROOT / "results/final_report_20261014_layout_audit.json"
    if destination.exists():
        raise FileExistsError("Preserve the accepted report receipt")
    docx = ROOT / "output/report/MicroC_实验报告_最终版_20261014.docx"
    pdf = ROOT / "output/pdf/MicroC_实验报告_最终版_20261014.pdf"
    reference = SUPPORT / "report_reference.docx"
    receipt = json.loads((SUPPORT / "report_edit_receipt_v3.json").read_text(encoding="utf-8"))
    assert digest(reference) == receipt["reference_sha256"]
    assert digest(docx) == receipt["output_sha256"] == digest(args.render_dir / "report.docx")
    assert digest(pdf) == digest(args.render_dir / "report.pdf")
    reviewed = sorted({int(x) for x in args.reviewed_pages.split(",")})
    page_count = len(PdfReader(pdf).pages)
    assert reviewed == list(range(1, page_count + 1)) == list(range(1, 29))
    pngs = sorted(args.render_dir.glob("page-*.png"))
    assert len(pngs) == page_count
    doc = Document(docx)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert not any(x in text for x in ("并不等于", "不等于", "■", "▪"))
    for prefix, count in (("表", 30), ("图", 14)):
        numbers = [int(m.group(1)) for p in doc.paragraphs
                   if p.style.name == "Caption" and (m := re.match(prefix + r"\s+(\d+)\s", p.text))]
        assert numbers == list(range(1, count + 1))
    assert len(doc.tables) == 30 and len(doc.inline_shapes) == 15
    assert all(t._tbl.find("w:tr/w:trPr/w:tblHeader", t._tbl.nsmap) is not None for t in doc.tables)
    assert not doc._element.xpath(".//w:trHeight")
    assert not doc._element.xpath(".//w:numPr")
    assert all(x._inline.docPr.get("descr") for x in doc.inline_shapes)
    with zipfile.ZipFile(docx) as final:
        assert all(hashlib.sha256(final.read(name)).hexdigest() == value
                   for name, value in receipt["preserved_package_parts"].items())
    result = {
        "accepted_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "hth810/micro-c", "pages": page_count,
        "reviewed_pages": reviewed,
        "visual_review": "All final pages inspected at original resolution; no clipping, black squares, broken tables or detached captions",
        "review_assignment": {"root": list(range(1, 15)), "final_report_visual_qa": list(range(15, 29))},
        "renderer": "Microsoft Word minimized isolated-copy PDF export; Poppler PNG 110 dpi",
        "bundled_renderer_fallback": "Bundled LibreOffice unavailable; failed renderer diagnosed before native Word export",
        "docx_sha256": digest(docx), "pdf_sha256": digest(pdf),
        "exact_export_source": True, "tables": 30, "figures": 14, "embedded_images": 15,
        "table_headers_complete": True, "image_alt_text_complete": True,
        "fixed_table_row_heights": 0, "automatic_bullet_paragraphs": 0,
        "preserved_package_parts": len(receipt["preserved_package_parts"]),
        "reference_sha256": digest(reference),
        "page_sha256": {p.name: digest(p) for p in pngs},
        "source_sha256": {"scripts/finalize_report_20261014.py": digest(Path(__file__))},
        "interpretation": "Visual/structural acceptance of the exact delivered report, separate from experiment conclusions"
    }
    destination.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("pages", "tables", "figures", "docx_sha256", "pdf_sha256")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
